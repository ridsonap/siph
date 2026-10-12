#!/usr/bin/env python3
"""
SIPH Sulsel Data Pipeline
Extracts IPH records from PDFs (text streams) and PNGs (Apple Vision OCR, macOS)
across 2023-2026 into a clean JSON dataset and prepares an optimized GeoJSON of
South Sulawesi regencies.

Validasi: tanda IPH dicek silang dengan kolom STATUS dan tanda andil komoditas.
Baris yang tetap bermasalah ditulis ke data/laporan_validasi.txt; hasil verifikasi
visual terhadap berkas sumber dicatat di data/koreksi_manual.json.
"""

import os
import re
import sys
import json
import zlib
import glob
import subprocess
import datetime
import urllib.request
from concurrent.futures import ProcessPoolExecutor
from PIL import Image

KAB_MAP = {
    '7301': 'Kab. Kepulauan Selayar',
    '7302': 'Kab. Bulukumba',
    '7303': 'Kab. Bantaeng',
    '7304': 'Kab. Jeneponto',
    '7305': 'Kab. Takalar',
    '7306': 'Kab. Gowa',
    '7307': 'Kab. Sinjai',
    '7308': 'Kab. Maros',
    '7309': 'Kab. Pangkajene Kepulauan',
    '7310': 'Kab. Barru',
    '7311': 'Kab. Bone',
    '7312': 'Kab. Soppeng',
    '7313': 'Kab. Wajo',
    '7314': 'Kab. Sidenreng Rappang',
    '7315': 'Kab. Pinrang',
    '7316': 'Kab. Enrekang',
    '7317': 'Kab. Luwu',
    '7318': 'Kab. Tana Toraja',
    '7322': 'Kab. Luwu Utara',
    '7325': 'Kab. Luwu Timur',
    '7326': 'Kab. Toraja Utara',
    '7371': 'Kota Makassar',
    '7372': 'Kota Parepare',
    '7373': 'Kota Palopo'
}

KAB_PATTERNS = [
    ('7322', r'luwu\s*utara'),
    ('7325', r'luwu\s*timur'),
    ('7317', r'\bluwu\b'),
    ('7326', r'toraja\s*utara'),
    ('7318', r'tana\s*toraja'),
    ('7301', r'selayar'),
    ('7302', r'bulukumba'),
    ('7303', r'bantaeng'),
    ('7304', r'jeneponto'),
    ('7305', r'takalar'),
    ('7306', r'\bgowa\b'),
    ('7307', r'sinjai'),
    ('7308', r'maros'),
    ('7309', r'pangkajene|pangkep'),
    ('7310', r'barru'),
    ('7311', r'\bbone\b'),
    ('7312', r'soppeng'),
    ('7313', r'\bwajo\b'),
    ('7314', r'sidenreng|sidrap'),
    ('7315', r'pinrang'),
    ('7316', r'enrekang'),
    ('7371', r'makassar'),
    ('7372', r'parepare'),
    ('7373', r'palopo')
]

MONTH_MAP = {
    'januari': 1, 'februari': 2, 'maret': 3, 'april': 4,
    'mei': 5, 'juni': 6, 'juli': 7, 'agustus': 8,
    'september': 9, 'oktober': 10, 'november': 11, 'desember': 12
}

MONTH_NAMES = [
    '', 'Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
    'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember'
]

def parse_filename(filepath):
    """Periode dari nama berkas, mis. '07.Juli M5 2026.pdf' atau '1. Januari_M2 2023.png'.
    Mengembalikan None bila bulan/minggu/tahun tidak dapat dikenali."""
    fn = os.path.basename(filepath).lower()
    m_found = next((num for name, num in MONTH_MAP.items() if name in fn), None)
    w_match = re.search(r'(?:^|[\s_.\-])m\s*([1-5])(?![0-9])', fn)
    y_match = re.search(r'(?<![0-9])(20[2-9][0-9])(?![0-9])', fn)
    if not (m_found and w_match and y_match):
        return None
    y_num, w_num = int(y_match.group(1)), int(w_match.group(1))
    return {
        'period_key': f"{y_num}-{m_found:02d}-W{w_num}",
        'year': y_num,
        'month': m_found,
        'week': w_num,
        'label': f"{MONTH_NAMES[m_found]} {y_num} (Minggu {w_num})",
        'filepath': filepath
    }

COMMODITY_CLEANUP = {
    'cabai rawit': 'Cabai Rawit',
    'cabairawit': 'Cabai Rawit',
    'cabai merah': 'Cabai Merah',
    'cabaimerah': 'Cabai Merah',
    'bawang merah': 'Bawang Merah',
    'bawangmerah': 'Bawang Merah',
    'bawang putih': 'Bawang Putih',
    'bawangputih': 'Bawang Putih',
    'daging ayam ras': 'Daging Ayam Ras',
    'dagingayamras': 'Daging Ayam Ras',
    'daging sapi': 'Daging Sapi',
    'dagingsapi': 'Daging Sapi',
    'telur ayam ras': 'Telur Ayam Ras',
    'telur ayamras': 'Telur Ayam Ras',
    'telurayamras': 'Telur Ayam Ras',
    'minyak goreng': 'Minyak Goreng',
    'minyakgoreng': 'Minyak Goreng',
    'udang basah': 'Udang Basah',
    'udangbasah': 'Udang Basah',
    'gula pasir': 'Gula Pasir',
    'gulapasir': 'Gula Pasir',
    'tepung terigu': 'Tepung Terigu',
    'tahu mentah': 'Tahu Mentah',
    'ikan kembung': 'Ikan Kembung',
    'ikan bandeng': 'Ikan Bandeng',
    'ikan tongkol': 'Ikan Tongkol',
    'ikan cakalang': 'Ikan Cakalang',
    'mie kering': 'Mie Kering Instan',
    'susu bubuk untuk balita': 'Susu Bubuk untuk Balita',
    'susu bubuk untuk bayi': 'Susu Bubuk untuk Bayi',
    'susu kental manis': 'Susu Kental Manis',
    'susu bubuk': 'Susu Bubuk',
    'bahan bakar': 'Bahan Bakar Rumah Tangga',
    'beras': 'Beras',
    'jeruk': 'Jeruk',
    'pisang': 'Pisang',
    'tempe': 'Tempe',
    'tomat': 'Tomat',
}

NUM = r'[-–—−]?\s*\d+(?:[\.,]\d+)?'
DEC_RE = re.compile(r'^[-–—−]?\s*\d+[\.,]\d+$')
IPH_RE = re.compile(r'^[-–—−]?\s*(?:\d{1,2}[\.,]\d+|\d)$')  # IPH boleh bulat satu digit ('0', '-1')
COMM_RE = re.compile(r'([A-Za-z][A-Za-z\s/\.\-]*?)\s*[\(\{\[]\s*(' + NUM + r')\s*[\)\}\]]')
STATUS_WORDS = {'naik': 'Naik', 'turun': 'Turun', 'stabil': 'Stabil'}
NO_FLUK = {'', '-', 'STABIL', 'NODATA', 'TETAP'}


def normalize_comm_name(raw):
    if not raw:
        return '-'
    clean_str = re.sub(r'\s+', ' ', re.sub(r'^[^\w]+|[^\w]+$', '', raw)).strip()
    s = re.sub(r'[^a-zA-Z\s]', '', clean_str).strip().lower().replace('cabal', 'cabai')
    s = re.sub(r'\s+', ' ', s)
    if s in COMMODITY_CLEANUP:
        return COMMODITY_CLEANUP[s]
    for k, v in COMMODITY_CLEANUP.items():  # urutan kunci: yang lebih spesifik lebih dulu
        if k in s:
            return v
    return clean_str.title()


def normalize_fluk(raw):
    t = re.sub(r'\s+', ' ', (raw or '')).strip(' _-–—|[]\t\n.,;')
    if t.upper() in NO_FLUK:
        return '-'
    return normalize_comm_name(t)


def clean_number(s):
    if not s:
        return None
    val = s
    for ch in '“”—–−=':
        val = val.replace(ch, '-')
    val = val.replace(',', '.').replace(' ', '').strip()
    try:
        return round(float(val), 4)
    except Exception:
        return None


def parse_komoditas(text):
    out = []
    for name, val in COMM_RE.findall(text or ''):
        name = name.strip(' ,;/.-')
        v = clean_number(val)
        if not name or v is None:
            continue
        n = normalize_comm_name(name)
        if any(w in n.lower() for w in ['kab', 'sulawesi', 'selatan', 'kode', 'provinsi']):
            continue
        out.append({'nama': n, 'kontribusi': v})
    return out


def kab_from_text(text):
    for code, pat in KAB_PATTERNS:
        if re.search(pat, text, re.I):
            return code
    return None


def build_record(code, iph, comms, fluk, cv, status_raw, kab_text=''):
    """Susun satu record + daftar catatan validasi (flags)."""
    flags = []
    status = 'Naik' if iph > 0 else ('Turun' if iph < 0 else 'Stabil')
    if status_raw:
        st = STATUS_WORDS.get(status_raw.strip().lower())
        if st and st != status:
            # Kolom STATUS (kata) lebih andal dari tanda minus hasil OCR
            flags.append(f'tanda IPH {iph} tidak sesuai status "{st}", dibalik')
            if (st == 'Turun' and iph > 0) or (st == 'Naik' and iph < 0):
                iph = -iph
            status = st
    if comms and iph != 0 and all((c['kontribusi'] > 0) != (iph > 0) for c in comms if c['kontribusi'] != 0) \
            and any(c['kontribusi'] != 0 for c in comms):
        if status_raw is None:
            # Tanpa kolom STATUS: tanda minus yang hilang terdeteksi dari tanda andil
            # (terverifikasi visual pada 20/20 kasus 2023-2024)
            flags.append(f'tanda IPH {iph} berlawanan dengan seluruh andil, dibalik')
            iph = -iph
            status = 'Naik' if iph > 0 else 'Turun'
        else:
            flags.append('tanda seluruh andil berlawanan dengan IPH')
    if abs(iph) > 15:
        flags.append(f'IPH tidak wajar ({iph})')
    kab_code = kab_from_text(kab_text) if kab_text else None
    if kab_code and kab_code != code:
        flags.append(f'kode {code} tidak cocok dengan nama kab ({kab_code})')
    return {
        'kode_kab': code,
        'nama_kab': KAB_MAP.get(code, f'Kab. {code}'),
        'iph': iph,
        'status': status,
        'komoditas': comms,
        'fluktuasi_tertinggi': fluk or '-',
        'cv': cv
    }, flags


def parse_row_tokens(tokens, has_status_col, iph_hint=None):
    """Urai token satu baris tabel: [kode, provinsi, kab, IPH, komoditas..., fluktuasi, CV, (status), (RPP)]."""
    code_i = next((i for i, t in enumerate(tokens) if re.match(r'^73\d{2}$', t)), None)
    head = tokens[:code_i] if code_i is not None else []
    code = tokens[code_i] if code_i is not None else None
    rest = tokens[code_i + 1:] if code_i is not None else tokens
    if code is None or code not in KAB_MAP:
        code = kab_from_text(' '.join(tokens)) or code
    if code is None or code not in KAB_MAP:
        return None, ['kode kabupaten tidak dikenali']

    if any('NODATA' in t.upper().replace(' ', '') for t in rest) and not any(DEC_RE.match(t) for t in rest):
        return None, []  # baris tanpa data
    # Lewati provinsi & nama kab (token alfabet tanpa angka) sampai sel IPH
    i = 0
    kab_parts = []
    while i < len(rest) and re.search(r'[A-Za-z]', rest[i]) and not re.search(r'\d', rest[i]):
        kab_parts.append(rest[i])
        i += 1
    if iph_hint is not None and (i >= len(rest) or not IPH_RE.match(rest[i])):
        rest = rest[:i] + [iph_hint] + rest[i:]
    if i >= len(rest):
        if not any(re.search(r'\d', t) for t in rest):
            return None, []  # baris tanpa angka sama sekali: tidak ada data
        return None, ['IPH tidak ditemukan']
    if rest[i].strip() in ('-', '—', '–'):
        return None, []  # tidak ada data minggu ini
    if not IPH_RE.match(rest[i]):
        return None, [f'IPH tidak terbaca (token: {rest[i][:40]!r})']
    iph = clean_number(rest[i])
    tail = rest[i + 1:]

    status_raw = None
    if has_status_col:
        for k in range(len(tail) - 1, -1, -1):
            if tail[k].strip().lower() in STATUS_WORDS:
                status_raw = tail[k]
                tail = tail[:k] + tail[k + 1:]
                break

    last_paren = max((k for k, t in enumerate(tail) if ')' in t or '(' in t), default=-1)
    comm_text = ' '.join(tail[:last_paren + 1])
    after = tail[last_paren + 1:]
    fluk_parts, cv = [], None
    for t in after:
        if DEC_RE.match(t) or re.match(r'^\d+$', t):
            cv = clean_number(t)
            break
        if t.strip() in ('-', '—', '–'):
            if fluk_parts:
                break
            continue
        fluk_parts.append(t)
    rec, flags = build_record(code, iph, parse_komoditas(comm_text), normalize_fluk(' '.join(fluk_parts)),
                              cv, status_raw, ' '.join(kab_parts))
    return rec, flags


# ---------------------------------------------------------------- PDF
def pdf_text_tokens(filepath):
    with open(filepath, 'rb') as f:
        content = f.read()
    for m in re.finditer(rb'/Length\s+(\d+).*?stream\r?\n', content, re.DOTALL):
        data = content[m.end():m.end() + int(m.group(1))]
        try:
            decompressed = zlib.decompress(data).decode('latin1', errors='ignore')
        except Exception:
            continue
        if 'BT' not in decompressed:
            continue
        tokens = []
        for bt in re.findall(r'BT\s*(.*?)\s*ET', decompressed, re.DOTALL):
            parts = re.findall(r'\(((?:\\.|[^\)])*)\)', bt)
            merged = ''.join(parts).replace(r'\(', '(').replace(r'\)', ')').strip()
            if merged:
                tokens.append(merged)
        if sum(1 for t in tokens if re.match(r'^73\d{2}$', t)) >= 5:
            yield tokens


def parse_pdf_file(filepath):
    meta = parse_filename(filepath)
    records, issues = [], []
    for tokens in pdf_text_tokens(filepath):
        has_status = any('STATUS' in t.upper() for t in tokens)
        idxs = [i for i, t in enumerate(tokens) if re.match(r'^73\d{2}$', t)]
        for n, i in enumerate(idxs):
            j = idxs[n + 1] if n + 1 < len(idxs) else len(tokens)
            rec, flags = parse_row_tokens(tokens[i:j], has_status)
            if rec:
                records.append(rec)
            if flags:
                issues.append((tokens[i], flags))
        if records:
            break
    return meta, records, issues


# ---------------------------------------------------------------- PNG (Apple Vision OCR)
OCR_CACHE = '.cache/ocr_vision.json'
OCR_TOOL = 'scripts/ocr_vision.swift'


def run_vision_ocr(png_files):
    cache = {}
    if os.path.exists(OCR_CACHE):
        with open(OCR_CACHE, encoding='utf-8') as f:
            cache = json.load(f)
    todo = [p for p in png_files if p not in cache]
    if todo:
        print(f"OCR Vision untuk {len(todo)} gambar...")
        os.makedirs(os.path.dirname(OCR_CACHE), exist_ok=True)
        tmp = OCR_CACHE + '.new'
        subprocess.check_call(['swift', OCR_TOOL, tmp] + todo)
        with open(tmp, encoding='utf-8') as f:
            cache.update(json.load(f))
        os.remove(tmp)
        with open(OCR_CACHE, 'w', encoding='utf-8') as f:
            json.dump(cache, f)
    return cache


def detect_row_bands(path, x0, x1):
    """Batas baris dari garis grid horizontal tabel (piksel gelap memanjang)."""
    im = Image.open(path).convert('L')
    W, H = im.size
    x0, x1 = max(0, int(x0)), min(W, int(x1))
    px = im.load()
    step = 3
    xs = range(x0, x1, step)
    need = 0.6 * len(xs)
    is_line = [sum(1 for x in xs if px[x, y] < 230) >= need for y in range(H)]
    runs, y = [], 0
    while y < H:
        if is_line[y]:
            s = y
            while y < H and is_line[y]:
                y += 1
            runs.append((s, y))
        y += 1
    edges = [(0, 0)] + runs + [(H, H)]  # baris teratas/terbawah bisa tanpa garis penutup
    bands = [(edges[k][1], edges[k + 1][0]) for k in range(len(edges) - 1) if edges[k + 1][0] - edges[k][1] > 6]
    return bands


def cluster_order(obs):
    """Urutan baca dalam satu baris: kelompokkan per kolom (tumpang-tindih x), lalu atas-ke-bawah."""
    cols = []
    for o in sorted(obs, key=lambda o: o['x']):
        a, b = o['x'], o['x'] + o['w']
        for c in cols:
            ov = min(b, c['b']) - max(a, c['a'])
            if ov > 0.5 * min(b - a, c['b'] - c['a']):
                c['items'].append(o)
                c['a'], c['b'] = min(a, c['a']), max(b, c['b'])
                break
        else:
            cols.append({'a': a, 'b': b, 'items': [o]})
    out = []
    for c in sorted(cols, key=lambda c: c['a']):
        out += sorted(c['items'], key=lambda o: o['y'])
    return out


def split_cell_text(text):
    """Pisahkan sel yang tergabung oleh OCR (mis. '7301 SULAWESI ...', '1,50 BAWANG(..)', 'CABAI RAWIT 0,28')."""
    t = text.replace('|', ' ').strip()
    out = []
    # Kode wilayah dengan sisa garis tabel di depannya: '7309, {7312, 17306
    m = re.match(r"^(?:\d{1,2}[\.\s]\s?)?[\'\"{\[(|!]?1?(73\d{2})(?:\s+(.*))?$", t)
    if m and m.group(1) in KAB_MAP:
        out.append(m.group(1))
        t = m.group(2) or ''
    # Desimal yang kehilangan koma: '-0 24' -> '-0.24'
    t = re.sub(r'^([-–—−]?\d)\s(\d{2,4})(?=\s|$)', r'\1.\2', t)
    m = re.match(r'^(' + NUM + r')\s+(.*\S.*)$', t)
    if m and IPH_RE.match(m.group(1).strip()) and not re.match(r'^\d', m.group(2)):
        out.append(m.group(1))
        t = m.group(2)
    m = re.match(r'^(.*?)\s+(naik|turun|stabil)$', t, re.I)
    status = None
    if m and '(' not in m.group(2):
        t, status = m.group(1), m.group(2)
    if '(' not in t:
        m = re.match(r'^(.*[A-Za-z].*?)\s+([-–—−]?\d+[\.,]\d+)(?:\s+([-–—−]?\d+[\.,]\d+))?$', t)
        if m:
            out += [m.group(1), m.group(2)] + ([m.group(3)] if m.group(3) else [])
            t = None
    if t:
        out.append(t)
    if status:
        out.append(status)
    return [x.strip() for x in out if x and x.strip()]


def iph_column_range(rows):
    """Perkiraan rentang x kolom IPH: sel pertama tiap baris yang diawali angka desimal."""
    xs = []
    for row in rows:
        for o in row:
            toks = split_cell_text(o['text'])
            if toks and re.match(r'^73\d{2}$', toks[0]):
                toks = toks[1:]
            if toks and DEC_RE.match(toks[0]):
                xs.append(o['x'])
                break
            if '(' in o['text']:
                break
    if not xs:
        return None
    xs.sort()
    x = xs[len(xs) // 2]
    return x - 25, x + 85


def iph_from_column(row, iph_cols):
    """Apakah baris memiliki sel berawalan angka desimal di dalam kolom IPH."""
    for o in row:
        if iph_cols[0] - 40 <= o['x'] <= iph_cols[1]:
            toks = split_cell_text(o['text'])
            if toks and re.match(r'^73\d{2}$', toks[0]):
                toks = toks[1:]
            if toks and IPH_RE.match(toks[0]):
                return True
        # sel kode+provinsi+kab+IPH yang tergabung oleh OCR
        if o['x'] < iph_cols[0] and o['x'] + o['w'] >= iph_cols[0] + 10 and re.search(r'\s[-–—−]?\d+(?:[\.,]\d+)?\s*$', o['text']):
            return True
    return False


def crop_key(filepath, b0, b1):
    return f"{filepath}|{int(b0)}|{int(b1)}"


def parse_png_file(filepath, ocr, crops=None):
    """Kembalikan (meta, records, issues, crop_requests)."""
    meta = parse_filename(filepath)
    obs = ocr['obs']
    has_status = any('STATUS' in o['text'].upper() for o in obs)
    x0 = min(o['x'] for o in obs)
    x1 = max(o['x'] + o['w'] for o in obs)
    bands = detect_row_bands(filepath, x0, x1)
    anchors = [o for o in obs if re.match(r"^\W?1?73\d{2}\b", o['text'].strip())]
    # Cadangan bila garis grid tidak terdeteksi: pita di antara titik tengah kode wilayah
    covered = sum(1 for a in anchors if any(b0 <= a['y'] + a['h'] / 2 <= b1 for b0, b1 in bands))
    if covered < 0.9 * len(anchors):
        ys = sorted(a['y'] + a['h'] / 2 for a in anchors)
        mids = [(ys[k] + ys[k + 1]) / 2 for k in range(len(ys) - 1)]
        edges = [ys[0] - 15] + mids + [ys[-1] + 15] if ys else []
        bands = list(zip(edges[:-1], edges[1:]))
    rows = []
    for b0, b1 in bands:
        row = cluster_order([o for o in obs if b0 <= o['y'] + o['h'] / 2 <= b1])
        if row:
            rows.append((b0, b1, row))
    iph_cols = iph_column_range([r for _, _, r in rows])
    records, issues, requests, seen = [], [], [], set()
    for b0, b1, row in rows:
        tokens = []
        for o in row:
            tokens += split_cell_text(o['text'])
        joined = ' '.join(tokens)
        if not re.search(r'\b73\d{2}\b', joined) and not kab_from_text(joined):
            continue
        if re.search(r'KODE|PROVINSI|KOMODITAS UTAMA|Kab/Kota', joined, re.I) and not re.search(r'\(', joined):
            continue
        rec, flags = parse_row_tokens(tokens, has_status)
        if rec and iph_cols and not iph_from_column(row, iph_cols):
            # angka yang terbaca sebagai IPH ternyata dari kolom lain (mis. CV saat sel IPH berisi '-')
            rec, flags = None, ['IPH tidak terbaca (bukan dari kolom IPH)']
        if not rec and flags and flags[0].startswith('IPH') and iph_cols:
            key = crop_key(filepath, b0, b1)
            hint = (crops or {}).get(key)
            if hint is None:
                requests.append((key, filepath, (iph_cols[0], b0, iph_cols[1], b1)))
            else:
                m = re.search(r'[-–—−]?(?:\d{1,2}[\.,]\d+|\d(?!\d))', hint.replace(' ', ''))
                if m:
                    rec, flags = parse_row_tokens(tokens, has_status, iph_hint=m.group(0))
                    flags = flags + [f'IPH dari OCR ulang potongan sel: {m.group(0)}']
                else:
                    rec, flags = None, []  # sel IPH kosong / '-': tidak ada data minggu ini
        if rec and rec['kode_kab'] in seen:
            flags = flags + ['baris ganda']
            rec = None
        if rec:
            seen.add(rec['kode_kab'])
            records.append(rec)
        if flags:
            issues.append((joined[:60], flags))
    return meta, records, issues, requests


CROP_CACHE = '.cache/ocr_crops.json'


def ocr_crops(requests):
    """OCR ulang potongan sel IPH (diperbesar 4x) untuk baris yang angkanya terlewat."""
    cache = {}
    if os.path.exists(CROP_CACHE):
        with open(CROP_CACHE, encoding='utf-8') as f:
            cache = json.load(f)
    todo = [r for r in requests if r[0] not in cache]
    if todo:
        tmpdir = '.cache/crops'
        os.makedirs(tmpdir, exist_ok=True)
        paths = {}
        for n, (key, filepath, (cx0, cy0, cx1, cy1)) in enumerate(todo):
            im = Image.open(filepath).convert('RGB')
            c = im.crop((max(0, int(cx0)), max(0, int(cy0)), min(im.width, int(cx1)), min(im.height, int(cy1))))
            c = c.resize((c.width * 4, c.height * 4), Image.LANCZOS)
            canvas = Image.new('RGB', (c.width + 80, c.height + 80), 'white')
            canvas.paste(c, (40, 40))
            p = os.path.join(tmpdir, f'crop_{n}.png')
            canvas.save(p)
            paths[p] = key
        out = tmpdir + '/hasil.json'
        subprocess.check_call(['swift', OCR_TOOL, out] + list(paths))
        with open(out, encoding='utf-8') as f:
            res = json.load(f)
        for p, key in paths.items():
            texts = sorted(res.get(p, {}).get('obs', []), key=lambda o: o['x'])
            cache[key] = ' '.join(o['text'] for o in texts)
        with open(CROP_CACHE, 'w', encoding='utf-8') as f:
            json.dump(cache, f)
    return cache


# ---------------------------------------------------------------- Excel / CSV
def read_table_rows(filepath):
    if filepath.lower().endswith('.csv'):
        import csv
        with open(filepath, encoding='utf-8-sig', newline='') as f:
            sample = f.read(4096)
            f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
            return [row for row in csv.reader(f, dialect)]
    from openpyxl import load_workbook
    wb = load_workbook(filepath, read_only=True, data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def parse_table_file(filepath):
    """Tabel IPH dari Excel/CSV dengan kolom seperti rilis BPS:
    KODE | ... | PERUBAHAN IPH | KOMODITAS ... | FLUKTUASI ... | NILAI CV | STATUS"""
    meta = parse_filename(filepath)
    rows = [[('' if c is None else str(c)).strip() for c in r] for r in read_table_rows(filepath)]
    hi = next((i for i, r in enumerate(rows) if any('KODE' in c.upper() for c in r)), None)
    if hi is None:
        return meta, [], [('berkas', ['header KODE tidak ditemukan'])]
    head = rows[hi]
    if hi + 1 < len(rows) and not any(re.match(r'^73\d{2}', c) for c in rows[hi + 1]):
        head = [f"{a} {b}" for a, b in zip(head, rows[hi + 1] + [''] * len(head))]
    head = [h.upper() for h in head]

    def col(*keys, exclude=()):
        return next((i for i, h in enumerate(head) if any(k in h for k in keys) and not any(x in h for x in exclude)), None)
    c_kode, c_iph = col('KODE'), col('PERUBAHAN', 'IPH', 'INDIKATOR', exclude=('KOMODITAS', 'CV'))
    c_kom, c_fluk = col('KOMODITAS'), col('FLUKTUASI', exclude=('CV', 'NILAI'))
    c_cv, c_status = col('CV'), col('STATUS')
    if c_kode is None or c_iph is None:
        return meta, [], [('berkas', ['kolom KODE / IPH tidak ditemukan'])]
    records, issues = [], []
    get = lambda r, i: r[i] if i is not None and i < len(r) else ''
    for r in rows[hi + 1:]:
        code = re.sub(r'\.0$', '', get(r, c_kode))
        if code not in KAB_MAP:
            continue
        iph = clean_number(get(r, c_iph))
        if iph is None:
            continue  # sel IPH kosong / '-'
        rec, flags = build_record(code, iph, parse_komoditas(get(r, c_kom)), normalize_fluk(get(r, c_fluk)),
                                  clean_number(get(r, c_cv)), get(r, c_status) or None)
        records.append(rec)
        if flags:
            issues.append((code, flags))
    return meta, records, issues


def apply_manual_corrections(meta, records, corrections):
    """Koreksi manual hasil verifikasi visual terhadap berkas sumber (data/koreksi_manual.json)."""
    fixes = corrections.get(meta['period_key'], {})
    by_code = {r['kode_kab']: r for r in records}
    for code, fix in fixes.items():
        if code.startswith('_'):
            continue
        r = by_code.get(code)
        if r is None:
            r = {'kode_kab': code, 'nama_kab': KAB_MAP[code], 'iph': None, 'status': None,
                 'komoditas': [], 'fluktuasi_tertinggi': '-', 'cv': None}
            records.append(r)
        for k, v in fix.items():
            if k == 'catatan':
                continue
            r[k] = v
        if r['iph'] is not None:
            r['status'] = 'Naik' if r['iph'] > 0 else ('Turun' if r['iph'] < 0 else 'Stabil')
    records.sort(key=lambda r: -(r['iph'] or 0))
    return records


def process_file(args):
    filepath, ocr, crops = args
    low = filepath.lower()
    if low.endswith('.pdf'):
        return parse_pdf_file(filepath) + ([],)
    if low.endswith(('.xlsx', '.csv')):
        return parse_table_file(filepath) + ([],)
    return parse_png_file(filepath, ocr, crops)


def parse_all(all_files, ocr_cache):
    """Lintasan 1: urai semua berkas; lintasan 2: OCR ulang sel IPH yang terlewat lalu urai ulang."""
    with ProcessPoolExecutor(max_workers=8) as ex:
        first = list(ex.map(process_file, [(f, ocr_cache.get(f), None) for f in all_files]))
    requests = [r for res in first for r in res[3]]
    if not requests:
        return first
    print(f"OCR ulang {len(requests)} sel IPH yang terlewat...")
    crops = ocr_crops(requests)
    with ProcessPoolExecutor(max_workers=8) as ex:
        return list(ex.map(process_file, [(f, ocr_cache.get(f), crops) for f in all_files]))


def fetch_and_optimize_geojson(output_path):
    url = 'https://raw.githubusercontent.com/TheMaggieSimpson/IndonesiaGeoJSON/main/kota-kabupaten.json'
    print(f"Downloading GeoJSON from {url}...")
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=15) as resp:
        d = json.loads(resp.read().decode())

    def round_coords(geom, precision=4):
        def _round(coords):
            if not coords: return []
            if isinstance(coords[0], (int, float)):
                return [round(coords[0], precision), round(coords[1], precision)]
            return [_round(c) for c in coords]
        return {'type': geom['type'], 'coordinates': _round(geom['coordinates'])}

    def simplify_ring(ring, tol=0.003):
        if len(ring) <= 4: return ring
        res = [ring[0]]
        for p in ring[1:-1]:
            if abs(p[0] - res[-1][0]) > tol or abs(p[1] - res[-1][1]) > tol:
                res.append(p)
        res.append(ring[-1])
        return res if len(res) >= 4 else ring

    def simplify_geom(geom):
        if geom['type'] == 'Polygon':
            return {'type': 'Polygon', 'coordinates': [simplify_ring(r) for r in geom['coordinates']]}
        elif geom['type'] == 'MultiPolygon':
            return {'type': 'MultiPolygon', 'coordinates': [[simplify_ring(r) for r in poly] for poly in geom['coordinates']]}
        return geom

    features = []
    for f in d.get('features', []):
        props = f.get('properties', {})
        if 'sulawesi selatan' in str(props).lower():
            code = str(props.get('CC_2', ''))
            name = props.get('NAME_2', '')
            type_kab = props.get('TYPE_2', 'Kabupaten')
            full_name = f"Kota {name}" if 'kota' in type_kab.lower() else f"Kab. {name}"
            # Standardize names to match KAB_MAP
            std_name = KAB_MAP.get(code, full_name)
            features.append({
                'type': 'Feature',
                'id': code,
                'properties': {
                    'id': code,
                    'code': code,
                    'name': std_name,
                    'kab_name': name
                },
                'geometry': simplify_geom(round_coords(f['geometry'], 4))
            })

    geojson = {'type': 'FeatureCollection', 'features': features}
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as out:
        json.dump(geojson, out, separators=(',', ':'))
    print(f"Saved optimized GeoJSON ({len(features)} features) -> {output_path} ({os.path.getsize(output_path)//1024} KB)")

def main():
    base_dir = 'Perkembangan IPH Sulsel'
    all_files = sorted(glob.glob(f'{base_dir}/**/*.*', recursive=True))
    all_files = [f for f in all_files if not os.path.basename(f).startswith(('.', '~$'))
                 and f.lower().endswith(('.png', '.jpg', '.jpeg', '.pdf', '.xlsx', '.csv'))]
    unnamed = [f for f in all_files if parse_filename(f) is None]
    if unnamed:
        print("⚠ Nama berkas berikut tidak memuat bulan, minggu (M1–M5), dan tahun, sehingga DILEWATI:")
        for f in unnamed:
            print(f"   - {f}")
        print("  Contoh nama yang benar: '11.November M2 2026.pdf'\n")
    all_files = [f for f in all_files if f not in unnamed]
    keys = {}
    for f in all_files:
        keys.setdefault(parse_filename(f)['period_key'], []).append(f)
    dup = {k: v for k, v in keys.items() if len(v) > 1}
    if dup:
        print("⚠ Lebih dari satu berkas untuk minggu yang sama (yang dipakai: berkas terakhir menurut abjad):")
        for k, v in dup.items():
            print(f"   - {k}: {', '.join(os.path.basename(x) for x in v)}")
        all_files = [f for f in all_files if f == keys[parse_filename(f)['period_key']][-1]]
    print(f"Ditemukan {len(all_files)} berkas sumber.")
    old_keys = set()
    if os.path.exists('data/iph_sulsel.json'):
        with open('data/iph_sulsel.json', encoding='utf-8') as f:
            old_keys = {p['period_key'] for p in json.load(f).get('periods', [])}

    pngs = [f for f in all_files if f.lower().endswith(('.png', '.jpg', '.jpeg'))]
    ocr_cache = run_vision_ocr(pngs) if pngs else {}
    corrections = {}
    if os.path.exists('data/koreksi_manual.json'):
        with open('data/koreksi_manual.json', encoding='utf-8') as f:
            corrections = json.load(f)

    results, report = [], []
    for meta, records, issues, _ in parse_all(all_files, ocr_cache):
        records = apply_manual_corrections(meta, records, corrections)
        fixed = set(corrections.get(meta['period_key'], {}))
        for where, flags in issues:
            code = re.search(r'\b(73\d{2})\b', where)
            auto = all(f.startswith(('IPH dari OCR ulang', 'tanda IPH')) for f in flags)
            status = ('DIKOREKSI MANUAL' if code and code.group(1) in fixed
                      else 'DIPERBAIKI OTOMATIS' if auto else 'PERLU DICEK')
            report.append(f"{meta['period_key']}\t{os.path.basename(meta['filepath'])}\t{status}\t{where}\t{'; '.join(flags)}")
        if records:
            results.append((meta, records))
            print(f"✓ [{meta['period_key']}] {os.path.basename(meta['filepath'])}: {len(records)} kab")

    with open('data/laporan_validasi.txt', 'w', encoding='utf-8') as f:
        f.write('periode\tberkas\tstatus\tbaris\tcatatan\n' + '\n'.join(report) + '\n')
    print(f"Catatan validasi: {len(report)} baris -> data/laporan_validasi.txt")

    # Sort results chronologically by period_key
    results.sort(key=lambda item: item[0]['period_key'])

    periods = []
    regencies_time_series = {code: {'name': name, 'series': []} for code, name in KAB_MAP.items()}
    commodity_stats = {}

    for meta, records in results:
        # calculate period stats
        iph_values = [r['iph'] for r in records if r['iph'] is not None]
        avg_iph = round(sum(iph_values) / len(iph_values), 2) if iph_values else 0
        naik_cnt = sum(1 for r in records if r['status'] == 'Naik')
        turun_cnt = sum(1 for r in records if r['status'] == 'Turun')
        stabil_cnt = sum(1 for r in records if r['status'] == 'Stabil')

        sorted_by_iph = sorted(records, key=lambda r: r['iph'] if r['iph'] is not None else 0)
        top_loser = sorted_by_iph[0] if sorted_by_iph else None
        top_gainer = sorted_by_iph[-1] if sorted_by_iph else None

        # Period commodity frequency
        freq = {}
        for r in records:
            if r['fluktuasi_tertinggi'] and r['fluktuasi_tertinggi'] != '-':
                fn = r['fluktuasi_tertinggi']
                freq[fn] = freq.get(fn, 0) + 1
                commodity_stats[fn] = commodity_stats.get(fn, 0) + 1

        top_threat = max(freq.items(), key=lambda x: x[1])[0] if freq else '-'

        period_entry = {
            'period_key': meta['period_key'],
            'year': meta['year'],
            'month': meta['month'],
            'week': meta['week'],
            'label': meta['label'],
            'avg_iph': avg_iph,
            'counts': {'naik': naik_cnt, 'turun': turun_cnt, 'stabil': stabil_cnt},
            'top_gainer': {'name': top_gainer['nama_kab'], 'iph': top_gainer['iph']} if top_gainer else None,
            'top_loser': {'name': top_loser['nama_kab'], 'iph': top_loser['iph']} if top_loser else None,
            'top_threat': top_threat,
            'records': records
        }
        periods.append(period_entry)

        # Append to regencies time series
        rec_map = {r['kode_kab']: r for r in records}
        for code, info in regencies_time_series.items():
            r = rec_map.get(code)
            info['series'].append({
                'period_key': meta['period_key'],
                'iph': r['iph'] if r else None,
                'status': r['status'] if r else None,
                'threat': r['fluktuasi_tertinggi'] if r else None
            })

    # Top overall commodities
    sorted_threats = sorted(commodity_stats.items(), key=lambda x: -x[1])[:10]

    output_dataset = {
        'meta': {
            'generated_at': datetime.date.today().isoformat(),
            'total_periods': len(periods),
            'years': sorted(list(set(p['year'] for p in periods))),
            'kab_count': len(KAB_MAP)
        },
        'periods': periods,
        'regencies_time_series': regencies_time_series,
        'top_overall_commodities': [{'name': name, 'count': cnt} for name, cnt in sorted_threats]
    }

    out_json = 'data/iph_sulsel.json'
    os.makedirs('data', exist_ok=True)
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(output_dataset, f, separators=(',', ':'))
    print(f"\n✓ Data tersimpan -> {out_json} ({os.path.getsize(out_json)//1024} KB)")
    new_keys = [p for p in periods if p['period_key'] not in old_keys]
    need_check = sum(1 for line in report if '\tPERLU DICEK\t' in line)
    print("\n================ RINGKASAN ================")
    print(f"Total rilis      : {len(periods)} minggu ({periods[0]['label']} – {periods[-1]['label']})")
    print(f"Rilis baru       : {len(new_keys)}" + (''.join(f"\n   + {p['label']}: {len(p['records'])} kab/kota" for p in new_keys)))
    if need_check:
        print(f"Perlu dicek      : {need_check} baris → lihat data/laporan_validasi.txt,")
        print("                   koreksi lewat data/koreksi_manual.json lalu jalankan ulang.")
    else:
        print("Perlu dicek      : tidak ada")
    print("===========================================")

    # Fetch and optimize GeoJSON (hanya bila belum ada, atau dengan --geo)
    if '--geo' in sys.argv or not os.path.exists('data/sulsel_geo.json'):
        fetch_and_optimize_geojson('data/sulsel_geo.json')

if __name__ == '__main__':
    main()
